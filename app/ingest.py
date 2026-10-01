"""Walk the corpus and turn every readable file into text chunks.

One bad file (unreadable, encrypted, corrupt, enormous) must never stop the
walk, so each file is parsed inside its own try/except and its own time limit.
Parsers return (chunks, images): chunks are (location, text); images are
(location, bytes) that still need reading by a vision model.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import signal
import sys
import time
from contextlib import contextmanager
from itertools import zip_longest
from pathlib import Path

MAX_CHARS = 1200        # target chunk size
FILE_TIMEOUT_S = 90     # one pathological file must not eat the whole index budget
MAX_READ = 64 << 20     # never read more than this much of one file
csv.field_size_limit(1 << 24)


def log(*a):
    print(*a, file=sys.stderr, flush=True)


@contextmanager
def time_limit(seconds):
    def boom(*_):
        raise TimeoutError(f"parse took longer than {seconds}s")

    try:
        old = signal.signal(signal.SIGALRM, boom)
    except ValueError:              # not the main thread: run without a limit
        yield
        return
    signal.alarm(seconds)
    try:
        yield
    finally:
        signal.alarm(0)
        signal.signal(signal.SIGALRM, old)


def windows(lines, max_chars=MAX_CHARS, overlap=2):
    """Pack lines into ~max_chars chunks sharing `overlap` lines with the next one.

    Yields (first_line_no, last_line_no, text).
    """
    flat = []
    for no, ln in enumerate(lines, 1):
        ln = ln.strip()
        while len(ln) > max_chars:          # one giant record must not become one giant chunk
            flat.append((no, ln[:max_chars]))
            ln = ln[max_chars - 200:]
        if ln:
            flat.append((no, ln))
    cur, size = [], 0
    for item in flat:
        if cur and size + len(item[1]) > max_chars:
            yield cur[0][0], cur[-1][0], "\n".join(t for _, t in cur)
            cur = cur[-overlap:]
            size = sum(len(t) + 1 for _, t in cur)
            if size > max_chars // 3:       # huge lines: overlapping them would bloat every chunk
                cur, size = [], 0
        cur.append(item)
        size += len(item[1]) + 1
    if cur:
        yield cur[0][0], cur[-1][0], "\n".join(t for _, t in cur)


def _read(path):
    with open(path, "rb") as f:
        raw = f.read(MAX_READ)
    if b"\0" in raw[:8192]:
        raise ValueError("binary file")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))              # 212.0 -> 212
    return str(v).strip()


def _table(rows, sheet=""):
    """One chunk per row. After the header row, every cell is labelled with its column name,
    so a row still says what it means once it is cut out of the table."""
    tag = f"[{sheet}] " if sheet else ""
    out, head = [], None
    for n, row in enumerate(rows, 1):
        cells = [c.strip() for c in row]
        if not any(cells):
            continue
        if head is None:
            if sum(map(bool, cells)) >= 2:
                head = cells
            out.append((f"{sheet} r{n}".strip(), tag + " | ".join(c for c in cells if c)))
        else:
            pairs = zip_longest(head, cells, fillvalue="")
            out.append((f"{sheet} r{n}".strip(), tag + " | ".join(f"{h}: {c}" if h else c for h, c in pairs if c)))
    return out


def parse_text(path):
    return [(f"lines {a}-{b}", t) for a, b, t in windows(_read(path).splitlines())], []


def parse_csv(path):
    text = _read(path)
    head = text.split("\n", 1)[0]
    delim = max(",;\t|", key=head.count) if any(d in head for d in ",;\t|") else ","
    return _table(csv.reader(io.StringIO(text), delimiter=delim)), []


def parse_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)   # ponytail: whole workbook in memory; read_only=True if sheets get huge
    try:
        out = []
        for ws in wb.worksheets:
            out += _table(([_cell(v) for v in r] for r in ws.iter_rows(values_only=True)), ws.title)
    finally:
        wb.close()
    return out, []


def parse_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    out, para = [], []

    def flush():
        out.extend((f"lines {a}-{b}", t) for a, b, t in windows(para))
        para.clear()

    for el in d.element.body.iterchildren():            # body order: paragraphs and tables interleaved
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para.append(Paragraph(el, d).text)
        elif tag == "tbl":
            flush()
            rows = []
            for r in Table(el, d).rows:
                cells, last = [], None
                for c in r.cells:                       # merged cells repeat; keep one
                    if c._tc is not last:
                        cells.append(c.text)
                    last = c._tc
                rows.append(cells)
            out += _table(rows, "table")
    flush()
    return out, []


def parse_pdf(path):
    from pypdf import PdfReader

    r = PdfReader(path)                 # AES-encrypted files can raise right here
    if r.is_encrypted:
        raise ValueError("encrypted")
    chunks, images = [], []
    for n, page in enumerate(r.pages, 1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:
            log(f"  {path.name} p{n}: text extraction failed: {e}")
            text = ""
        if text:
            chunks += [(f"p{n} lines {a}-{b}", t) for a, b, t in windows(text.splitlines())]
        else:                           # scanned page: hand its embedded picture to the vision model
            try:
                images += [(f"p{n}", im.data) for im in page.images]
            except Exception as e:
                log(f"  {path.name} p{n}: no text and no readable image: {e}")
    return chunks, images


def parse_image(path):
    with open(path, "rb") as f:
        return [], [("image", f.read(MAX_READ))]


PARSERS = {
    ".pdf": parse_pdf, ".docx": parse_docx, ".xlsx": parse_xlsx, ".csv": parse_csv,
    ".txt": parse_text, ".log": parse_text, ".py": parse_text,
    **dict.fromkeys((".png", ".jpg", ".jpeg", ".tif", ".tiff"), parse_image),
}
KIND = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".csv": "csv",
        ".txt": "text", ".log": "text", ".py": "text",
        **dict.fromkeys((".png", ".jpg", ".jpeg", ".tif", ".tiff"), "image")}

# A withdrawn / superseded document must never be cited. Only the filename and the title
# block count: a current revision's history may well say "revision 1 is withdrawn".
BAD_NAME = re.compile(r"withdrawn|superseded|obsolete", re.I)
BAD_TITLE = re.compile(r"withdrawn|superseded by|obsolete", re.I)
TITLED = {".pdf", ".docx", ".txt"}


def status(rel, ext, first_text):
    title = " ".join(first_text.split("\n", 3)[:3]) if ext in TITLED else ""
    return "withdrawn" if BAD_NAME.search(rel) or BAD_TITLE.search(title) else "ok"


def walk(corpus):
    for dirpath, dirs, names in os.walk(corpus, onerror=lambda e: log("cannot list", e)):
        dirs.sort()
        for n in sorted(names):
            yield Path(dirpath, n)


def ingest(corpus, ocr=None, ocr_budget_s=240):
    """Parse every file under corpus. `ocr(bytes) -> str` reads pictures; without it they are indexed by name only.

    Returns (chunks, files): chunks are dicts {file, loc, kind, text, status};
    files is the manifest of what was indexed and what was skipped, and why.
    """
    corpus = Path(corpus)
    chunks, files = [], []
    t0 = time.monotonic()
    for path in walk(corpus):
        rel = path.relative_to(corpus).as_posix()
        ext = path.suffix.lower()
        if ext not in PARSERS:
            files.append({"path": rel, "status": "skipped", "why": "no parser for this type"})
            continue
        try:
            with time_limit(FILE_TIMEOUT_S):
                texts, images = PARSERS[ext](path)
                texts = list(texts)
                for loc, data in images:
                    read = ""
                    if ocr and time.monotonic() - t0 < ocr_budget_s:
                        try:
                            read = ocr(data).strip()
                        except Exception as e:
                            log(f"  {rel}: vision model failed: {e}")
                    if read or KIND[ext] == "image":
                        texts.append((loc, read))   # a picture stays retrievable by name even if unread
            if not texts:
                files.append({"path": rel, "status": "skipped", "why": "no text found"})
                continue
            st = status(rel, ext, texts[0][1])
            chunks += [{"file": rel, "loc": loc, "kind": KIND[ext], "text": t, "status": st} for loc, t in texts]
            files.append({"path": rel, "status": st, "chunks": len(texts)})
        except Exception as e:      # PermissionError, encrypted, corrupt, TimeoutError, ...
            why = f"{type(e).__name__}: {e}"[:200]
            files.append({"path": rel, "status": "skipped", "why": why})
            log(f"skip {rel}: {why}")
    return chunks, files


def save(directory, chunks, files):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    tmp = directory / "chunks.jsonl.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(c, ensure_ascii=False) + "\n")
    os.replace(tmp, directory / "chunks.jsonl")         # atomic: a reader never sees half an index
    (directory / "files.json").write_text(json.dumps(files, indent=1), encoding="utf-8")


def load(directory):
    path = Path(directory) / "chunks.jsonl"
    if not path.exists():
        return None
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]
